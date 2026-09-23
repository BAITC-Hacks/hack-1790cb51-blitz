"""Build long, paired ATLAS fixtures and an independent answer key.

Run with the Codex bundled Python runtime. Does not call OpenAI or modify app code.
"""

import copy
import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, KeepTogether, PageBreak

from content import TOPICS

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "hackalem-long-documents"
QA = ROOT / ".test-data" / "fixture-qa"

DETAILS = [
    "Ответственный специалист фиксирует исходные данные, дату выполнения и принятое решение в рабочем реестре. При расхождении сведений он запрашивает уточнение у владельца исходного документа до завершения операции.",
    "Результат передается руководителю подразделения вместе с перечнем использованных документов. Неустраненные замечания отражаются отдельно, чтобы получатель мог отличить завершенную работу от вопроса, требующего решения.",
    "Материалы готовятся по каждому филиалу отдельно и затем сводятся в общий пакет. Повторное получение одного и того же документа не является основанием для двойного учета соответствующей операции.",
    "Исполнитель сохраняет обоснование решения и согласованную версию документа. Изменение исходных данных после согласования требует новой записи с указанием причины и даты изменения.",
    "Рабочий реестр содержит владельца вопроса, состояние исполнения и ссылку на подтверждающий материал. Завершенная запись остается доступной для последующей проверки и не заменяется новой версией без сохранения истории.",
    "Перед передачей результата смежному подразделению исполнитель проверяет комплектность материалов. Если часть сведений отсутствует, в сопроводительной записи указываются недостающие сведения и ответственное за них лицо.",
    "Результат рассматривается руководителем на ежемесячном рабочем совещании. Решение по спорному вопросу фиксируется отдельно от исходного обращения, чтобы сохранить последовательность его рассмотрения.",
    "Каждый отчетный период обрабатывается независимо от предыдущего. Сопоставление с прошлым периодом выполняется по одинаковому составу объектов, а обнаруженные различия сопровождаются пояснением причин.",
]


def para(core, section_index, index):
    return core.rstrip(".") + ". " + DETAILS[(section_index * 3 + index) % len(DETAILS)]


def paraphrase(text):
    replacements = [
        ("Формирование", "Подготовка"), ("Подготовка", "Разработка"),
        ("Ведение", "Поддержание в актуальном состоянии"),
        ("Контроль", "Проверка"), ("Проведение", "Организация"),
        ("Согласование", "Рассмотрение и согласование"),
        ("Оценка", "Анализ"), ("Разработка", "Подготовка"),
        ("Анализ", "Оценка"),
        ("Хранение", "Обеспечение сохранности"), ("Регистрация", "Учет"),
    ]
    for before, after in replacements:
        if text.startswith(before):
            text = after + text[len(before):]
            break
    return text.replace("исходные данные", "первоначальные сведения").replace(
        "подтверждающий материал", "документальное подтверждение"
    ).replace("руководителю подразделения", "начальнику подразделения")


def dataset(topic_key, config):
    before = []
    for s, (department, heading, duties) in enumerate(config["sections"]):
        before.append({"key": f"s{s}", "department": department, "heading": heading,
            "items": [{"uid": f"{topic_key}-{s}-{i}", "origins": [f"{topic_key}-{s}-{i}"],
                "text": para(core, s, i), "core": core, "changes": []}
                for i, core in enumerate(duties)]})
    sections = copy.deepcopy(before)
    events = []
    def event(kind, old, new, explanation):
        events.append({"type": kind, "before_ids": old, "after_ids": new, "expectation": explanation})
    def ids(section):
        return [i["uid"] for i in section["items"]]

    sections[0]["department"] = config["renamed"]
    sections[0]["heading"] = "Планирование и подготовка решений"
    event("rename", ids(before[0]), ids(sections[0]), "Название подразделения и заголовок изменены; обязанности сохранены. Не считать весь раздел потерянным.")
    for index in (0, 2, 4):
        item = sections[0]["items"][index]
        item["text"] = paraphrase(item["text"])
        item["changes"].append("paraphrase")
        event("paraphrase", item["origins"], [item["uid"]], "Смысл обязанности сохранен при замене формулировки.")

    # A compound duty becomes two separate duties owned by separate new units.
    split_source = sections[1]["items"][2]
    halves = []
    for index in range(2):
        item = copy.deepcopy(split_source)
        item["uid"] += f"-part{index + 1}"
        item["core"] = config["split_function"][index]
        item["text"] = para(item["core"], 1, 2)
        item["changes"] = ["function_split"]
        halves.append(item)
    first = {"key": "split-a", "department": config["split"][0], "heading": "Планирование и контроль", "items": sections[1]["items"][:2] + [halves[0]]}
    second = {"key": "split-b", "department": config["split"][1], "heading": "Планирование и контроль", "items": [halves[1]] + sections[1]["items"][3:]}
    second["items"][-1]["text"] = paraphrase(second["items"][-1]["text"])
    event("section_split", ids(before[1]), ids(first) + ids(second), "Один раздел разделен между двумя подразделениями. Заголовки обоих новых разделов совпадают, владельцы различаются.")
    event("function_split", split_source["origins"], [i["uid"] for i in halves], "Одна составная функция имеет двух преемников. Ни один отдельный пункт не покрывает исходную обязанность целиком.")

    sections[2]["heading"] = "Материалы и записи отчетного периода"
    event("section_move", ids(before[2]), ids(sections[2]), "Раздел перенесен в конец документа и переименован; исходные обязанности этого раздела не изменены.")
    changed = sections[3]["items"][0]
    changed["text"] = changed["text"].replace("5 рабочих дней", "2 рабочих дней")
    changed["changes"].append("deadline_change")
    event("deadline_change", changed["origins"], [changed["uid"]], "Срок сокращен с 5 до 2 рабочих дней. Тема функции сохранена, требование изменилось.")
    changed = sections[3]["items"][1]
    changed["text"] = para(config["role_change"], 3, 1)
    changed["changes"].append("authority_change")
    event("authority_change", changed["origins"], [changed["uid"]], "Утверждение заменено подготовкой рекомендации; добавлено явное отрицание права утверждать. Не считать полномочия полностью сохраненными.")
    removed = sections[3]["items"][-2:]
    sections[3]["items"] = sections[3]["items"][:-2]
    event("functions_removed", [i["uid"] for i in removed], [], "Две обязанности удалены без назначения преемника. Найти обе возможные потери.")
    event("section_removed", ids(before[4]), [], "Раздел обучения и все его обязанности отсутствуют в новой редакции. Это контрольный случай полной потери раздела.")

    for s in (5, 6):
        sections[s]["department"] = config["merged"]
        sections[s]["heading"] = "Регистрация и контроль"
    for item in sections[6]["items"][::2]:
        item["text"] = paraphrase(item["text"])
    event("sections_merged", ids(before[5]) + ids(before[6]), ids(sections[5]) + ids(sections[6]), "Два старых подразделения объединены в один центр. В новой редакции два одноименных подраздела относятся к разным предметам работы, это не автоматическое дублирование.")

    moved = sections[0]["items"].pop()
    moved["text"] = paraphrase(moved["text"])
    sections[7]["items"].append(moved)
    event("function_transfer", moved["origins"], [moved["uid"]], "Функция перенесена в другое подразделение и перефразирована; это передача, не утрата.")
    moved = sections[3]["items"].pop()
    sections[2]["items"].insert(0, moved)
    event("subsection_transfer", moved["origins"], [moved["uid"]], "Пункт перенесен из контрольного подраздела в другой раздел без изменения текста.")

    duplicate = copy.deepcopy(sections[0]["items"][1])
    duplicate["uid"] += "-copy"
    sections[7]["items"].append(duplicate)
    event("exact_duplication", duplicate["origins"], [sections[0]["items"][1]["uid"], duplicate["uid"]], "Одинаковая обязанность закреплена за двумя подразделениями без разделения ролей; потенциальное дублирование.")
    duplicate = copy.deepcopy(sections[0]["items"][3])
    duplicate["uid"] += "-reworded-copy"
    duplicate["text"] = paraphrase(duplicate["text"])
    sections[7]["items"].append(duplicate)
    event("semantic_duplication", duplicate["origins"], [sections[0]["items"][3]["uid"], duplicate["uid"]], "Одна обязанность продублирована у другого владельца с перефразированием.")
    for index, core in enumerate(config["new"]):
        item = {"uid": f"{topic_key}-new-{index}", "origins": [], "core": core, "text": para(core, 7, index), "changes": ["new"]}
        sections[7]["items"].append(item)
        event("new_function", [], [item["uid"]], "Новый пункт без прямого предшественника; проверить возможное пересечение с близкими функциями.")
    if topic_key == "operations":
        original = before[3]["items"][3]
        moved_audit = next(i for i in sections[3]["items"] if i["uid"] == original["uid"])
        sections[3]["items"].remove(moved_audit)
        sections[0]["items"].append(moved_audit)
        event("execution_control_conflict", [original["uid"], before[0]["items"][2]["uid"]], [moved_audit["uid"], before[0]["items"][2]["uid"]], "В дирекции снабжения совмещены выбор поставщика и независимый аудит закупок. Это индикатор конфликта исполнения и контроля.")
    after = [sections[3], sections[0], sections[5], second, sections[7], sections[6], first, sections[2]]
    for phase, collection in (("before", before), ("after", after)):
        for s_index, section in enumerate(collection, 1):
            for i_index, item in enumerate(section["items"], 1):
                item["number"] = f"{s_index}.{i_index}"
                item["location"] = f"Раздел {s_index} / {section['department']} / пункт {item['number']}"
    old = {i["uid"]: i for s in before for i in s["items"]}
    new = {i["uid"]: i for s in after for i in s["items"]}
    for item in events:
        item["before"] = [{"location": old[x]["location"], "text": old[x]["text"]} for x in item["before_ids"]]
        item["after"] = [{"location": new[x]["location"], "text": new[x]["text"]} for x in item["after_ids"] if x in new]
    return {"key": topic_key, "config": config, "before": before, "after": after, "events": events}


def docx_file(data, phase):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.1
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.add_paragraph(data["config"]["title" if phase == "before" else "after_title"], "Title")
    doc.add_paragraph(data["config"]["company"], "Subtitle")
    doc.add_paragraph("Редакция до реорганизации" if phase == "before" else "Редакция после реорганизации")
    doc.add_paragraph(data["config"]["purpose"])
    doc.add_paragraph("Правила работы с материалами", "Heading 1")
    doc.add_paragraph("Каждый раздел закрепляет функции за указанным подразделением. Руководитель распределяет задачи внутри подразделения, а исполнитель фиксирует результат в предусмотренном рабочем реестре. Передача сведений смежному подразделению сама по себе не означает передачу права принятия окончательного решения.")
    for number, group in enumerate(data[phase], 1):
        doc.add_paragraph(f"{number} {group['heading']}", "Heading 1")
        owner = doc.add_paragraph("Подразделение: " + group["department"])
        owner.paragraph_format.keep_with_next = True
        for index, item in enumerate(group["items"]):
            if index in (0, len(group["items"]) // 2):
                doc.add_paragraph("Основные обязанности" if index == 0 else "Порядок выполнения", "Heading 2")
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = True
            p.add_run(item["number"] + ". ").bold = True
            p.add_run(item["text"])
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Страница ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = "Команда Blitz"
    doc.core_properties.title = data["config"]["title" if phase == "before" else "after_title"]
    doc.save(OUT / f"finance-{phase}.docx")


def pdf_file(data, phase):
    fonts = Path("C:/Windows/Fonts")
    pdfmetrics.registerFont(TTFont("AtlasArial", str(fonts / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("AtlasArialBold", str(fonts / "arialbd.ttf")))
    body = ParagraphStyle("BodyRU", fontName="AtlasArial", fontSize=11, leading=15, spaceAfter=9, alignment=TA_LEFT)
    heading = ParagraphStyle("HeadRU", parent=body, fontName="AtlasArialBold", fontSize=14, leading=18, spaceBefore=15, spaceAfter=9, keepWithNext=True)
    title = ParagraphStyle("TitleRU", parent=heading, fontSize=23, leading=28, spaceBefore=0, spaceAfter=13)
    small = ParagraphStyle("SmallRU", parent=body, fontSize=10, leading=13)
    story = [Paragraph(escape(data["config"]["title" if phase == "before" else "after_title"]), title),
        Paragraph(escape(data["config"]["company"]), body),
        Paragraph("Редакция до реорганизации" if phase == "before" else "Редакция после реорганизации", small),
        Paragraph(escape(data["config"]["purpose"]), body)]
    for number, group in enumerate(data[phase], 1):
        if phase == "after" and number == len(data[phase]):
            story.append(PageBreak())
        # Keep the heading, owner and first duty together across page breaks.
        first = group["items"][0]
        story.append(KeepTogether([
            Paragraph(f"{number} {escape(group['heading'])}", heading),
            Paragraph("Подразделение: " + escape(group["department"]), body),
            Paragraph(escape(first["number"] + ". " + first["text"]), body),
        ]))
        for item in group["items"][1:]:
            story.append(KeepTogether([Paragraph(escape(item["number"] + ". " + item["text"]), body)]))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("AtlasArial", 9)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawRightString(A4[0] - 48, 26, f"Страница {doc.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(OUT / f"legal-{phase}.pdf"), pagesize=A4,
        rightMargin=48, leftMargin=48, topMargin=42, bottomMargin=44,
        title=data["config"]["title" if phase == "before" else "after_title"], author="Команда Blitz").build(story, onFirstPage=footer, onLaterPages=footer)


def answer_key(all_data):
    lines = ["# Контрольная карта изменений", "", "**Не загружайте эту карту в ATLAS вместе с исходными документами.** Она содержит ответы для ручной проверки.", "",
        "Все организации, обязанности и сроки вымышлены. Юридическая пара не является изложением законодательства или правовой консультацией. Изменения внутри документов намеренные; модель может обоснованно сгруппировать некоторые замечания иначе.", "",
        "Проверяйте смысл, владельца, срок и полномочие отдельно. Переименование не равно потере; одинаковый подзаголовок не равен дублированию. Отсутствие пункта и отсутствие семантического преемника тоже не одно и то же.", ""]
    for data in all_data:
        lines += [f"## {data['config']['company']}", "", f"Тема: {data['config']['title']}. До: {sum(len(s['items']) for s in data['before'])} функций. После: {sum(len(s['items']) for s in data['after'])} функций.", ""]
        for index, event in enumerate(data["events"], 1):
            lines += [f"### {index} {event['type']}", "", event["expectation"], "", "До:", ""]
            lines += [f"- {r['location']}: {r['text'].split('. ')[0]}." for r in event["before"]] or ["- Отсутствует."]
            lines += ["", "После:", ""]
            lines += [f"- {r['location']}: {r['text'].split('. ')[0]}." for r in event["after"]] or ["- Отсутствует."]
            lines += [""]
    (OUT / "EXPECTED-CHANGES.md").write_text("\n".join(lines), encoding="utf-8")
    (OUT / "expected-changes.json").write_text(json.dumps(all_data, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    data = [dataset(key, config) for key, config in TOPICS.items()]
    for phase in ("before", "after"):
        docx_file(data[0], phase)
        pdf_file(data[1], phase)
    answer_key(data)
    (QA / "datasets.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({d["key"]: {p: sum(len(s["items"]) for s in d[p]) for p in ("before", "after")} for d in data}))


if __name__ == "__main__":
    main()
