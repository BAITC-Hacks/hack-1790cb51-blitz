"""Offline fixture verification and distribution archive; never calls an AI API."""

import json
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "hackalem-long-documents"
sys.path.insert(0, str(ROOT / "backend"))
from app.parser import extract  # noqa: E402


def normalize(text):
    return re.sub(r"\s+", " ", text).strip()


datasets = json.loads((OUT / "expected-changes.json").read_text(encoding="utf-8"))
files = []
for data in datasets:
    extension = {"finance": "docx", "legal": "pdf", "operations": "xlsx"}[data["key"]]
    for phase in ("before", "after"):
        path = OUT / f"{data['key']}-{phase}.{extension}"
        segments, warnings = extract(path.name, path.read_bytes())
        assert 0 < len(segments) <= 450, (path.name, len(segments))
        extracted = normalize(" ".join(s["text"] for s in segments))
        items = [item for section in data[phase] for item in section["items"]]
        for item in items:
            assert normalize(item["text"]) in extracted, (path.name, item["number"])
        for section in data[phase]:
            assert section["department"] in extracted
        assert "sk-proj-" not in extracted
        assert "EXPECTED-CHANGES" not in extracted
        if extension == "pdf":
            assert len(PdfReader(path).pages) == 6
        if extension == "xlsx":
            book = load_workbook(path)
            assert len(book.worksheets) == 3
            assert sum(s.max_row - 5 for s in book) == len(items)
            assert all(s.freeze_panes == "B6" for s in book)
            assert all(len(s.tables) == 1 for s in book)
            assert not any(c.data_type == "e" for s in book for row in s for c in row)
            book.close()
        files.append(path)
        print(f"PASS {path.name}: {len(items)} duties; {len(segments)} source fragments")
    for event in data["events"]:
        if event["type"] == "exact_duplication":
            assert event["after"][0]["text"] == event["after"][1]["text"]
        if event["type"] == "semantic_duplication":
            assert event["after"][0]["text"] != event["after"][1]["text"]
        if event["type"] == "paraphrase":
            assert event["before"][0]["text"] != event["after"][0]["text"]

with zipfile.ZipFile(OUT / "hackalem-test-documents.zip", "w", zipfile.ZIP_DEFLATED) as archive:
    for path in files + [OUT / "README.md", OUT / "EXPECTED-CHANGES.md", OUT / "expected-changes.json"]:
        archive.write(path, path.name)
print("PASS archive: six input documents, instructions, and separate answer keys")
